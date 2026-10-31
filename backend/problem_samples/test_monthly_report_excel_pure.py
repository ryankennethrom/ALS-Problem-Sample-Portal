"""Standalone regression tests for safe XLSX downloads (no Django database required)."""
import ast
import io
import unittest
from pathlib import Path
from types import SimpleNamespace
import json
import re
from openpyxl import load_workbook

SOURCE=Path(__file__).with_name('monthly_report_portal.py')
functions=[n for n in ast.parse(SOURCE.read_text()).body if isinstance(n,ast.FunctionDef)
           and n.name in {'_safe_text','_excel_value','reportable_column','render_excel'}]
namespace={'io':io,'json':json,'re':re}
exec(compile(ast.fix_missing_locations(ast.Module(body=functions,type_ignores=[])),str(SOURCE),'exec'),namespace)

class XLSXDownloadPureTests(unittest.TestCase):
    def test_all_tickets_preserve_order_and_header(self):
        snapshot={'total':2,'workflow_counts':{'Disposed':1,'Follow Up Required':1},
                  'columns':[{'id':'a','label':'Customer'},{'id':'b','label':'Batch'}],
                  'tickets':[{'number':10,'status':'Disposed','details':[{'id':'b','label':'Batch','value':'B1'},{'id':'a','label':'Customer','value':'C1'}]},
                             {'number':20,'status':'Follow Up Required','details':[{'id':'b','label':'Batch','value':'B2'}]}]}
        book=load_workbook(io.BytesIO(namespace['render_excel'](snapshot)),read_only=True)
        rows=list(book['Tickets'].values)
        self.assertEqual(rows[1][2:],('Customer','Batch'))
        self.assertEqual(rows[2][2:],('C1','B1'))
        self.assertEqual(rows[3][2:],(None,'B2'))
        self.assertEqual(rows[2][0],10)
        self.assertEqual(rows[1][:2], ('Ticket #', 'Current Workflow'))
        self.assertNotIn('Created', rows[1])
        self.assertNotIn('Brand Matched Via', rows[1])
        self.assertEqual(book.sheetnames, ['Summary', 'Tickets'])
        self.assertEqual(rows[3][0],20)

    def test_explanatory_ticket_row_is_removed(self):
        snapshot = {'columns': [{'id':'batch', 'label':'Batch'}],
                    'tickets': [{'number': 27, 'status': 'Disposed',
                                 'details': [{'id':'batch','label':'Batch','value':'Lot-27'}]}]}
        book = load_workbook(io.BytesIO(namespace['render_excel'](snapshot)))
        ws = book['Tickets']
        self.assertEqual(ws['A1'].value, 'ALS TICKET TRACKER — MONTHLY BRAND REPORT')
        self.assertEqual([ws.cell(2, i).value for i in range(1, 4)],
                         ['Ticket #', 'Current Workflow', 'Batch'])
        self.assertEqual([ws.cell(3, i).value for i in range(1, 4)],
                         [27, 'Disposed', 'Lot-27'])
        self.assertNotIn('All matching tickets appear below',
                         ' '.join(str(c.value or '') for row in ws for c in row))
        self.assertEqual(ws.auto_filter.ref, 'A2:C3')
        self.assertEqual(ws.print_title_rows, '$2:$2')

    def test_legacy_report_blocks_sensitive_columns(self):
        legacy={'tickets':[{'number':13,'details':[{'label':'Password','value':'123456'},
                                                   {'label':'Batch','value':'A'}]}]}
        workbook=load_workbook(io.BytesIO(namespace['render_excel'](legacy)),read_only=True)
        self.assertEqual(workbook['Tickets']['C2'].value,'Batch')
        self.assertNotIn('Password',str(list(workbook['Tickets'].values)))

    def test_untrusted_spreadsheet_formulas_never_evaluated(self):
        payload={'columns':[{'id':'a','label':'Comment'}],
                 'tickets':[{'number':1,'details':[{'id':'a','label':'Comment','value':'=2+2'}]}]}
        book=load_workbook(io.BytesIO(namespace['render_excel'](payload)),read_only=True,data_only=False)
        cell=book['Tickets']['C3']
        self.assertEqual(cell.value,"'=2+2")
        self.assertNotEqual(cell.data_type,'f')

    def test_header_filters_wrap_freeze_and_row_height(self):
        snapshot={'columns':[{'id':'note','label':'Notes'}], 'total':1,
                  'tickets':[{'number':7,'status':'Follow Up Required','details':[
                      {'id':'note','label':'Notes','value':'A long note ' * 60}]}]}
        book=load_workbook(io.BytesIO(namespace['render_excel'](snapshot)))
        ws=book['Tickets']
        self.assertEqual(ws.auto_filter.ref, 'A2:C3')
        self.assertEqual(ws.freeze_panes, 'C3')
        self.assertTrue(ws['C3'].alignment.wrap_text)
        self.assertGreater(ws.row_dimensions[3].height, 27)

    def test_old_saved_snapshots_hide_obsolete_fixed_columns(self):
        # Existing report links may contain these snapshot keys, but must not
        # output their old fixed columns in newly downloaded workbooks.
        snapshot={'columns':[{'id':'a','label':'Brand'},{'id':'b','label':'Created'},
                             {'id':'c','label':'Brand Matched Via'}],
                  'tickets':[{'number':5,'status':'Disposed','created_at':'2026-10-29',
                              'matched_sources':['Distributor Brand'], 'details':[
                                  {'id':'a','label':'Brand','value':'Northstar'},
                                  {'id':'b','label':'Created','value':'2026-10-29'},
                                  {'id':'c','label':'Brand Matched Via','value':'Distributor Brand'}]}]}
        book=load_workbook(io.BytesIO(namespace['render_excel'](snapshot)),read_only=True)
        ws=book['Tickets']
        self.assertEqual(tuple(ws.values)[1], ('Ticket #', 'Current Workflow', 'Brand'))
        self.assertEqual(tuple(ws.values)[2], (5, 'Disposed', 'Northstar'))

    def test_empty_selected_columns_produce_two_fixed_columns(self):
        snapshot={'columns':[],'tickets':[{'number':11,'status':'Disposed',
                                          'created_at':'2026-10-29','matched_sources':['Ticket Brand']}],
                  'total':1}
        book=load_workbook(io.BytesIO(namespace['render_excel'](snapshot)))
        ws=book['Tickets']
        self.assertEqual(ws['A2'].value, 'Ticket #')
        self.assertEqual(ws['B2'].value, 'Current Workflow')
        self.assertEqual(ws['C2'].value, None)
        self.assertEqual(ws.auto_filter.ref, 'A2:B3')
        self.assertEqual(ws.freeze_panes, 'C3')
        self.assertEqual(ws['A3'].value, 11)

    def test_summary_tab_explains_where_ticket_rows_are(self):
        workbook=load_workbook(io.BytesIO(namespace['render_excel']({'total':2,'tickets':[{'number':1}]})), read_only=True)
        self.assertEqual(workbook.sheetnames, ['Summary','Tickets'])
        self.assertIn('Open the Tickets worksheet tab', str(workbook['Summary']['B6'].value))
        self.assertEqual(workbook['Tickets']['A3'].value,1)

if __name__=='__main__':unittest.main()
