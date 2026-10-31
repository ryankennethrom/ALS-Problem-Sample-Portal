"""Django integration tests: run after installing backend requirements."""
from datetime import date, datetime, timedelta
import io
from openpyxl import load_workbook
from zoneinfo import ZoneInfo
from django.test import TestCase, override_settings
from django.contrib.auth.models import User
from accounts.models import UserProfile
from django.utils import timezone
from rest_framework.test import APIClient

from accounts.models import Lab
from .models import ProblemColumn, ProblemSample, ProblemTable, SystemEmail
from .system_email_service import create_monthly_report, queue_manual_monthly_report


@override_settings(FRONTEND_URL='https://tracker.example.test')
class PublicMonthlyReportTests(TestCase):
    def setUp(self):
        self.lab, _ = Lab.objects.get_or_create(code='EDM', defaults={'name':'Edmonton'})
        self.table = ProblemTable.objects.create(name='ALS Tickets')
        self.column = ProblemColumn.objects.create(table=self.table, name='Brand', field_key='brand', column_type='brand')
        self.row = ProblemSample.objects.create(lab=self.lab, table=self.table, problem_number=990,
            current_workflow='Follow Up Required', custom_values={'brand':'Northstar'},
            problem_type='Missing paperwork', issue_description='Please confirm batch ID.')
        self.rule = {'pattern':'Northstar', 'match_type':'contains', 'sources':['ticket'],
                     'cc_emails':['manager@example.com'], 'id':'2a661ae5-2d4c-47c8-9f8f-b1a77ce0ed12'}
        import uuid
        self.mail, _ = queue_manual_monthly_report(self.column,self.rule,
            request_identifier=uuid.uuid4(),lab=self.lab)
        self.client = APIClient() # anonymous link recipients
        self.url = f'/api/public/monthly-reports/{self.mail.public_token}/xlsx/'

    def test_anonymous_xlsx_has_eligible_fields_and_private_headers(self):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertIn('.xlsx', response['Content-Disposition'])
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(book.sheetnames, ['Summary', 'Tickets'])
        self.assertEqual(book['Tickets']['A3'].value, 990)
        self.assertEqual(book['Tickets']['B3'].value, 'Follow Up Required')
        self.assertEqual(response['Cache-Control'], 'private, no-store, max-age=0')

    def test_email_contains_link_not_ticket_detail(self):
        self.assertIn(f'/reports/monthly/{self.mail.public_token}', self.mail.body)
        self.assertIn('open the Tickets tab', self.mail.body)
        self.assertNotIn('Please confirm batch ID.', self.mail.body)
        self.assertNotIn('Ticket #990', self.mail.body)

    def test_selected_columns_frozen_and_excluded_fields_omitted(self):
        batch = ProblemColumn.objects.create(table=self.table, name='Batch', field_key='batch', column_type='text')
        hidden = ProblemColumn.objects.create(table=self.table, name='Secret', field_key='secret-key', column_type='text')
        self.row.custom_values = {'brand':'Northstar','batch':'AB-17','secret-key':'not permitted'}
        self.row.save()
        rule = dict(self.rule, report_column_ids=[str(batch.pk)])
        snapshot = create_monthly_report(self.column, rule, date(2026,10,1), date(2026,11,1), [self.row], lab=self.lab)['report_snapshot']
        self.assertEqual([c['label'] for c in snapshot['columns']], ['Batch'])
        self.assertNotIn('not permitted',str(snapshot))
        self.mail.report_snapshot = snapshot
        self.mail.save(update_fields=['report_snapshot'])
        response = self.client.get(self.url)
        book = load_workbook(io.BytesIO(response.content), read_only=True)
        self.assertEqual(book['Tickets']['C2'].value,'Batch')
        self.assertEqual(book['Tickets']['C3'].value,'AB-17')
        self.assertIsNone(book['Tickets']['D2'].value)

    def test_report_snapshot_does_not_change_when_ticket_updates(self):
        old = load_workbook(io.BytesIO(self.client.get(self.url).content), read_only=True)['Tickets']['B3'].value
        self.row.current_workflow = 'Disposed'
        self.row.save()
        self.assertEqual(load_workbook(io.BytesIO(self.client.get(self.url).content), read_only=True)['Tickets']['B3'].value, old)

    def test_expired_revoked_and_missing_links_denied(self):
        SystemEmail.objects.filter(pk=self.mail.pk).update(public_expires_at=timezone.now()-timedelta(seconds=1))
        self.assertEqual(self.client.get(self.url).status_code,404)
        SystemEmail.objects.filter(pk=self.mail.pk).update(public_expires_at=timezone.now()+timedelta(days=30),public_revoked_at=timezone.now())
        self.assertEqual(self.client.get(self.url).status_code,404)
        SystemEmail.objects.filter(pk=self.mail.pk).update(public_revoked_at=None)
        self.assertEqual(self.client.get(self.url).status_code,200)

    def test_only_admin_can_revoke_report_link(self):
        normal = User.objects.create_user(username='ordinary', email='ordinary@example.com')
        UserProfile.objects.create(user=normal, role=UserProfile.ROLE_CUSTOMER_SERVICE, lab=self.lab)
        self.client.force_authenticate(normal)
        endpoint = f'/api/system-emails/{self.mail.pk}/revoke-report/'
        self.assertEqual(self.client.post(endpoint).status_code,403)
        admin = User.objects.create_user(username='admin2', email='admin2@alsglobal.com')
        UserProfile.objects.create(user=admin, is_admin=True, lab=self.lab)
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.post(endpoint).status_code,200)
        self.client.force_authenticate(user=None)
        self.assertEqual(self.client.get(self.url).status_code,404)

    def test_invalid_identifier_cannot_read_report(self):
        self.assertEqual(self.client.get('/api/public/monthly-reports/00000000-0000-0000-0000-000000000000/xlsx/').status_code,404)
